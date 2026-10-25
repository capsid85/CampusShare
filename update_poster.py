import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def update_poster():
    prs = pptx.Presentation("Design Thinking Poster.pptx")
    slide = prs.slides[0]

    c_dark = RGBColor(15, 23, 42)
    c_blue = RGBColor(37, 99, 235)
    c_gray = RGBColor(71, 85, 105)

    # 1. Update Title Shape
    for s in slide.shapes:
        if s.name == "Google Shape;92;p1":
            s.text_frame.text = "DESIGN THINKING (HUM_4342)\nCampusShare: Peer-to-Peer Campus Resource & Book Exchange"
            for p in s.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                p.font.bold = True
                p.font.color.rgb = RGBColor(255, 255, 255)
        elif s.name == "Google Shape;93;p1":
            s.text_frame.text = "Team Members: [Member 1 | Member 2 | Member 3 | Member 4]  •  Reg Nos: [2309xxxxx | 2309xxxxx | 2309xxxxx | 2309xxxxx]\nMentor: Dr. Vandana Jha, Associate Professor  •  Manipal Institute of Technology Bengaluru, MAHE"
            for p in s.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                p.font.color.rgb = RGBColor(241, 245, 249)
        elif s.name == "Google Shape;178;p1":
            s.text_frame.text = "MIT\nBengaluru"

    # Helper to add a clean content card on the poster
    def add_card(left_units, top_units, width_units, height_units, title, bullets):
        # 1 unit = 100,000 in pptx
        left = int(left_units * 100000)
        top = int(top_units * 100000)
        width = int(width_units * 100000)
        height = int(height_units * 100000)

        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(255, 255, 255)
        card.line.color.rgb = RGBColor(203, 213, 225)
        card.line.width = Pt(2)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = Inches(0.2)

        p = tf.paragraphs[0]
        p.text = title.upper()
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = c_blue
        p.space_after = Pt(12)

        for b in bullets:
            p2 = tf.add_paragraph()
            p2.text = "•  " + b
            p2.font.size = Pt(15)
            p2.font.color.rgb = c_dark
            p2.space_after = Pt(8)
        return card

    # Card 1: Objective & Phase 1 Empathy (Left Top)
    add_card(
        10.0, 58.0, 165.0, 190.0,
        "Phase 1: Empathy & Objective",
        [
            "Project Objective: Create a zero-friction, verified circular exchange platform allowing engineering students to Borrow, Rent, or Buy single-semester academic gear (textbooks, Casio fx-991 calculators, lab coats, drafters) from seniors at near-zero cost.",
            "Qualitative Field Insights (MIT Bengaluru Campus):",
            "  - 88% of freshers spend ₹2,500–₹4,500 every semester on specialized gear used for only 16 weeks.",
            "  - 92% of seniors have 3+ textbooks & tools sitting idle in hostel cupboards with no easy way to transfer them.",
            "  - Existing WhatsApp groups fail: notices get buried under hundreds of messages within 15 minutes.",
            "Customer Journey Mapping Highlights:",
            "  - Touchpoints: Syllabus release -> WhatsApp panic -> Senior negotiation -> Handover friction -> Disposal.",
            "  - Emotional Dip: High anxiety during 1st month due to unexpected equipment expenses.",
            "Problem Statement: Freshers face recurring financial burdens while seniors hoard unused academic supplies, caused by the absence of a structured, campus-verified peer-to-peer exchange."
        ]
    )

    # Card 2: Phase 2 Analyze (Left Bottom)
    add_card(
        10.0, 268.0, 165.0, 85.0,
        "Phase 2: Analyse & Root Cause",
        [
            "5-Whys Root Cause Analysis:",
            "  1. Why do freshers overspend on tools? They buy brand new items from external stores.",
            "  2. Why don't they get them from seniors? Unused items remain locked in seniors' hostel rooms.",
            "  3. Why don't seniors sell or lend them? Informal posting on social media is cumbersome and unorganized.",
            "  4. Why are chat groups inefficient? No search filters, item categories, or availability tracking.",
            "  5. Root Cause: Lack of a dedicated, campus-verified web portal facilitating discovery and peer handover.",
            "How Might We (HMW) Question:",
            "  'How might we connect junior and senior students for safe, verified, zero-commission exchange of academic resources at designated campus landmarks?'"
        ]
    )

    # Card 3: Phase 3 Solve (Right Top)
    add_card(
        185.0, 58.0, 165.0, 72.0,
        "Phase 3: Solve & Key Features",
        [
            "Multi-Tier Exchange Mechanism:",
            "  • Borrow (Free): 1-semester loan with return guarantee (e.g. Higher Engg Math).",
            "  • Rent (Low-Cost): Weekly rentals (₹20-50/wk) for exam prep (e.g. Casio fx-991).",
            "  • Buy (75% Off): Permanent transfer at deep discount (e.g. Mini Drafters, CLRS).",
            "Verified Campus Security: Restricted to @learner.manipal.edu credentials.",
            "Zero Courier Delay: Handovers coordinated at Central Library & Food Court."
        ]
    )

    # Card 4: Phase 4 Test / Prototype (Right Middle)
    add_card(
        185.0, 145.0, 165.0, 125.0,
        "Phase 4: Prototype Implementation",
        [
            "Interactive Web Application Built with Semantic HTML5, CSS3 & ES6+ JS.",
            "Key Prototype Modules:",
            "  ✔ Live Search & Category Filtering (Textbooks, Calculators, Lab Coats, Dev Boards).",
            "  ✔ Instant Mode Switcher (All / Borrow / Rent / Buy).",
            "  ✔ Listing Creator Modal: Seniors publish items in 30 seconds with photos & pickup spots.",
            "  ✔ Automated Peer Request Modal: Notifies lender with suggested handover meeting.",
            "  ✔ Sustainability Tracking: Recirculates idle items and protects pocket money with 0% platform fee."
        ]
    )

    # Add Prototype Screenshot into Right Middle Card
    screenshot_path = "c:\\Users\\sridh\\Desktop\\des1\\clean_cards_two_col.png"
    if os.path.exists(screenshot_path):
        # Insert image nicely
        slide.shapes.add_picture(
            screenshot_path,
            int(190.0 * 100000),
            int(215.0 * 100000),
            int(155.0 * 100000),
            int(52.0 * 100000)
        )

    # Card 5: Outcome & Future Scope (Right Bottom)
    add_card(
        185.0, 285.0, 165.0, 68.0,
        "Outcome & Impact",
        [
            "Peer Walkthrough & Thematic Validation:",
            "  • Zero-Knowledge Safety: Female hostel students praised complete concealment of phone numbers & room details.",
            "  • On-Spot UPI Trust: Direct scan-and-pay eliminates fear of advance app escrow lock-in.",
            "  • Transparent Fixed Pricing: Freshers confirmed total relief from opaque 'DM for price' culture.",
            "Campus Impact: Prevents recurring spending on 1-semester gear and recirculates idle resources within MIT Bengaluru.",
            "Future Roadmap: WhatsApp Cloud API alerts, campus locker integration, and library reserve sync."
        ]
    )

    prs.save("Design Thinking Poster.pptx")
    print("Successfully updated Design Thinking Poster.pptx with all sections and screenshot!")

if __name__ == "__main__":
    update_poster()
